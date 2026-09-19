import streamlit as st
import requests
import json
from docx import Document
from pypdf import PdfReader
import io

WEBHOOK_URL = "https://dustinjang7.app.n8n.cloud/webhook/kanvas-listing-generate"

st.set_page_config(
    page_title="Kanvas Marketplace Listing Generator",
    page_icon="📦",
    layout="wide"
)

st.title("📦 Kanvas Product Brief Listing Generator")
st.markdown("Upload a product brief (`.pdf` or `.docx`), parse its contents, and send it to n8n to generate marketplace listings.")

uploaded_file = st.file_uploader("Upload Product Brief", type=["pdf", "docx"])

def extract_text(file):
    text = ""
    if file.name.endswith(".pdf"):
        reader = PdfReader(file)
        for page in reader.pages:
            text += page.extract_text() or ""
    elif file.name.endswith(".docx"):
        doc = Document(file)
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
    return text

def create_docx(data):
    doc = Document()
    doc.add_heading("Generated Marketplace Listing", level=1)
    
    if isinstance(data, dict):
        if "target_market" in data:
            doc.add_paragraph(f"Target Market: {data['target_market']}")
            
        if "title_en" in data:
            doc.add_heading("Title (English)", level=2)
            doc.add_paragraph(data["title_en"])
            
        if data.get("title_localized"):
            doc.add_heading("Title (Localized)", level=2)
            doc.add_paragraph(data["title_localized"])
            
        if "bullet_points" in data and isinstance(data["bullet_points"], list):
            doc.add_heading("Bullet Points", level=2)
            for bullet in data["bullet_points"]:
                doc.add_paragraph(bullet, style="List Bullet")
                
        if "description" in data:
            doc.add_heading("Product Description", level=2)
            doc.add_paragraph(data["description"])
    else:
        doc.add_paragraph(str(data))
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

if uploaded_file:
    if st.button("🚀 Process Brief & Generate Listing", type="primary"):
        with st.spinner("Extracting text and calling n8n workflow..."):
            # Step 1: Extract text
            extracted_text = extract_text(uploaded_file)
            
            # Step 2: Payload
            payload = {
                "filename": uploaded_file.name,
                "brief_text": extracted_text
            }
            
            # Step 3: Webhook call
            try:
                response = requests.post(WEBHOOK_URL, json=payload, timeout=60)
                if response.status_code == 200:
                    st.success("Listing generated successfully!")
                    
                    try:
                        result_json = response.json()
                    except Exception:
                        result_json = {"raw_response": response.text}
                    
                    col1, col2 = st.columns(2)
                    
                    with col1:
                        st.subheader("📄 Generated JSON")
                        st.json(result_json)
                        
                    with col2:
                        st.subheader("💾 Export Options")
                        docx_bytes = create_docx(result_json)
                        
                        st.download_button(
                            label="⬇️ Download Output as .docx",
                            data=docx_bytes,
                            file_name=f"generated_listing_{uploaded_file.name.split('.')[0]}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )
                else:
                    st.error(f"Error from webhook ({response.status_code}): {response.text}")
            except requests.exceptions.RequestException as e:
                st.error(f"Failed to connect to n8n Webhook: {e}")